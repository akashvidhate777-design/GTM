"""Builds the campaign workbook (.xlsx):

  * "Leads & Sequences" — every lead + Lead Score/Grade + all 6 messages +
    Status dropdown + sequence-tracking columns.
  * "Campaign Report"  — live COUNTIF-driven funnel (updates as Status changes).
  * "Config"           — scoring weights and cadence, for reference.
"""
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .models import ScoredLead

HEADERS = [
    "First Name", "Last Name", "Title", "Company", "Email", "# Employees",
    "Industry", "Person Linkedin Url", "Website", "City", "State", "Country",
    "Annual Revenue",
    # -- added by the agent --
    "Lead Score", "Lead Grade", "ICP Tier", "Persona", "News Signal",
    "LinkedIn Connection Message", "LinkedIn Follow-Up 1", "LinkedIn Follow-Up 2",
    "Email 1 Subject", "Email 1 Body", "Email 2 Subject", "Email 2 Body",
    "Email 3 Subject", "Email 3 Body",
    "Status", "Sequence Step", "Last Touch Date", "Next Action Date", "Notes",
]

STATUS_OPTIONS = [
    "Draft", "Approved", "Sent - Step 1", "Sent - Step 2", "Sent - Step 3",
    "Replied", "Bounced", "Do Not Contact",
]

MESSAGE_COLS = {
    "LinkedIn Connection Message", "LinkedIn Follow-Up 1", "LinkedIn Follow-Up 2",
    "Email 1 Body", "Email 2 Body", "Email 3 Body",
}

GRADE_FILLS = {
    "A": PatternFill("solid", start_color="C6EFCE"),
    "B": PatternFill("solid", start_color="FFEB9C"),
    "C": PatternFill("solid", start_color="FFE4C4"),
    "D": PatternFill("solid", start_color="F2F2F2"),
}

HEADER_FILL = PatternFill("solid", start_color="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _col(name: str) -> str:
    return get_column_letter(HEADERS.index(name) + 1)


def _leads_sheet(wb: Workbook, scored: list[ScoredLead]):
    ws = wb.active
    ws.title = "Leads & Sequences"

    for c, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "E2"

    for r, item in enumerate(scored, 2):
        l, s, q = item.lead, item.score, item.sequence
        values = {
            "First Name": l.first_name, "Last Name": l.last_name,
            "Title": l.title, "Company": l.company, "Email": l.email,
            "# Employees": l.employees, "Industry": l.industry,
            "Person Linkedin Url": l.person_linkedin, "Website": l.website,
            "City": l.city, "State": l.state, "Country": l.country,
            "Annual Revenue": l.annual_revenue,
            "Lead Score": s.total, "Lead Grade": s.grade,
            "ICP Tier": s.icp_tier, "Persona": s.persona,
            "News Signal": s.news_signal,
            "LinkedIn Connection Message": q.linkedin_connection,
            "LinkedIn Follow-Up 1": q.linkedin_followup_1,
            "LinkedIn Follow-Up 2": q.linkedin_followup_2,
            "Email 1 Subject": q.email_1_subject, "Email 1 Body": q.email_1_body,
            "Email 2 Subject": q.email_2_subject, "Email 2 Body": q.email_2_body,
            "Email 3 Subject": q.email_3_subject, "Email 3 Body": q.email_3_body,
            "Status": item.status, "Sequence Step": 0,
            "Last Touch Date": "", "Next Action Date": "", "Notes": "",
        }
        for c, h in enumerate(HEADERS, 1):
            cell = ws.cell(row=r, column=c, value=values.get(h))
            if h in MESSAGE_COLS:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        grade_cell = ws.cell(row=r, column=HEADERS.index("Lead Grade") + 1)
        fill = GRADE_FILLS.get(str(grade_cell.value)[:1])
        if fill:
            grade_cell.fill = fill
            ws.cell(row=r, column=HEADERS.index("Lead Score") + 1).fill = fill

    n = len(scored) + 1
    dv = DataValidation(
        type="list", formula1=f'"{",".join(STATUS_OPTIONS)}"', allow_blank=True,
        showDropDown=False,
    )
    ws.add_data_validation(dv)
    dv.add(f"{_col('Status')}2:{_col('Status')}{max(n, 2)}")

    step_dv = DataValidation(type="list", formula1='"0,1,2,3"', allow_blank=True)
    ws.add_data_validation(step_dv)
    step_dv.add(f"{_col('Sequence Step')}2:{_col('Sequence Step')}{max(n, 2)}")

    widths = {
        "Title": 34, "Company": 28, "Email": 30, "Person Linkedin Url": 34,
        "Website": 24, "Persona": 30, "ICP Tier": 24, "News Signal": 30,
        "LinkedIn Connection Message": 50, "LinkedIn Follow-Up 1": 50,
        "LinkedIn Follow-Up 2": 50, "Email 1 Subject": 34, "Email 1 Body": 60,
        "Email 2 Subject": 34, "Email 2 Body": 60, "Email 3 Subject": 34,
        "Email 3 Body": 60, "Status": 16, "Notes": 30,
    }
    for c, h in enumerate(HEADERS, 1):
        ws.column_dimensions[get_column_letter(c)].width = widths.get(h, 15)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(n, 2)}"
    return ws


def _report_sheet(wb: Workbook, scored: list[ScoredLead]):
    ws = wb.create_sheet("Campaign Report")
    L = "'Leads & Sequences'"
    grade_c, status_c = _col("Lead Grade"), _col("Status")
    score_c, ind_c = _col("Lead Score"), _col("Industry")

    def header(row, text):
        cell = ws.cell(row=row, column=1, value=text)
        cell.font = Font(bold=True, size=12, color="1F3864")

    ws.cell(row=1, column=1, value="Campaign Report — AI on ERP for Retail Ecosystems").font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value=f"Generated {date.today().isoformat()} · live formulas — refresh as Status changes")

    header(4, "Pipeline")
    rows = [
        ("Total leads", f"=COUNTA({L}!{_col('Email')}2:{_col('Email')}10000)"),
        ("Average lead score", f"=ROUND(AVERAGE({L}!{score_c}2:{score_c}10000),1)"),
        ("Grade A - Hot (≥90)", f'=COUNTIF({L}!{grade_c}2:{grade_c}10000,"A*")'),
        ("Grade B - Warm (75-89)", f'=COUNTIF({L}!{grade_c}2:{grade_c}10000,"B*")'),
        ("Grade C - Nurture (55-74)", f'=COUNTIF({L}!{grade_c}2:{grade_c}10000,"C*")'),
        ("Grade D - Low (<55)", f'=COUNTIF({L}!{grade_c}2:{grade_c}10000,"D*")'),
    ]
    for i, (label, formula) in enumerate(rows, 5):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=formula)

    header(12, "Sequence funnel")
    funnel = [("Draft", "Draft"), ("Approved", "Approved"),
              ("Sent - Step 1", "Sent - Step 1"), ("Sent - Step 2", "Sent - Step 2"),
              ("Sent - Step 3", "Sent - Step 3"), ("Replied", "Replied"),
              ("Bounced", "Bounced"), ("Do Not Contact", "Do Not Contact")]
    for i, (label, status) in enumerate(funnel, 13):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=f'=COUNTIF({L}!{status_c}2:{status_c}10000,"{status}")')
    ws.cell(row=21, column=1, value="Total sent (any step)")
    ws.cell(row=21, column=2, value=f'=COUNTIF({L}!{status_c}2:{status_c}10000,"Sent*")+B18')
    ws.cell(row=22, column=1, value="Reply rate")
    ws.cell(row=22, column=2, value='=IFERROR(B18/B21,0)').number_format = "0.0%"

    header(24, "Leads by industry")
    industries = sorted({s.lead.industry for s in scored if s.lead.industry})
    for i, ind in enumerate(industries, 25):
        ws.cell(row=i, column=1, value=ind)
        ws.cell(row=i, column=2, value=f'=COUNTIF({L}!{ind_c}2:{ind_c}10000,"{ind}")')

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 14
    return ws


def _config_sheet(wb: Workbook):
    ws = wb.create_sheet("Config")
    lines = [
        ("Scoring weights", ""),
        ("ICP title tier (C-suite/Founder 40, VP 30, Director/Head 24, Manager 12, other 5)", 40),
        ("Function relevance (Ops/SC/Tech/Finance high)", 15),
        ("Industry fit (retail ecosystem core 25, adjacent 16, other 6)", 25),
        ("Company scale (employees, mid-market sweet spot)", 10),
        ("Annual revenue", 5),
        ("Data completeness (email + LinkedIn)", 5),
        ("News signal bonus (n8n web-search enrichment)", 5),
        ("", ""),
        ("Grades", ""),
        ("A - Hot", ">= 90"),
        ("B - Warm", "75-89"),
        ("C - Nurture", "55-74"),
        ("D - Low (parked, no sequence)", "< 55"),
        ("", ""),
        ("Cadence", ""),
        ("Step 1 (Day 0)", "LinkedIn connection + Email 1"),
        ("Step 2 (Day 3)", "LinkedIn follow-up 1 + Email 2"),
        ("Step 3 (Day 7)", "LinkedIn follow-up 2 + Email 3"),
        ("Stop conditions", "Replied / Bounced / Do Not Contact"),
        ("Daily send cap", 30),
    ]
    for r, (a, b) in enumerate(lines, 1):
        ws.cell(row=r, column=1, value=a)
        ws.cell(row=r, column=2, value=b)
        if b == "":
            ws.cell(row=r, column=1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["B"].width = 34
    return ws


def build_workbook(scored: list[ScoredLead], path: str):
    scored = sorted(scored, key=lambda s: -s.score.total)
    wb = Workbook()
    _leads_sheet(wb, scored)
    _report_sheet(wb, scored)
    _config_sheet(wb)
    wb.save(path)
    return path
